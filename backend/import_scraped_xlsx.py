"""One-shot import of the latest scraped competitor xlsx into our products
collection. Wipes the existing catalogue first (per user choice a-ii from the
session).

Image strategy: keep the smartpet.ge CDN URLs as `image`. They're high-res and
stable; if any go down, the admin can re-upload through the product editor.
"""

import asyncio
import re
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path

from motor.motor_asyncio import AsyncIOMotorClient
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))

import os
from dotenv import load_dotenv

load_dotenv(Path(__file__).parent / ".env")
MONGO_URL = os.environ["MONGO_URL"]
DB_NAME = os.environ["DB_NAME"]


def _slugify(text: str) -> str:
    base = re.sub(r"[^a-z0-9]+", "-", (text or "").lower()).strip("-")
    return base or f"product-{uuid.uuid4().hex[:8]}"


def _cell_str(v):
    return str(v).strip() if v is not None else ""


def _cell_float(v):
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", "."))
    except (ValueError, TypeError):
        return None


def _cell_bool(v):
    if v is None:
        return False
    if isinstance(v, bool):
        return v
    return str(v).strip().lower() in {"true", "1", "yes", "y"}


def _cell_tags(v):
    if not v:
        return []
    return [p.strip() for p in re.split(r"[;,]", str(v)) if p.strip()]


async def main(xlsx_path: Path) -> None:
    print(f"Reading {xlsx_path}")
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    ws = wb["Products"] if "Products" in wb.sheetnames else wb.worksheets[0]

    rows_iter = ws.iter_rows(values_only=True)
    headers = [_cell_str(c).lower() for c in next(rows_iter)]
    docs = []
    now_iso = datetime.now(timezone.utc).isoformat()

    for raw in rows_iter:
        if not any(_cell_str(v) for v in raw):
            continue
        d = dict(zip(headers, raw))
        name = _cell_str(d.get("name"))
        if not name:
            continue
        brand = _cell_str(d.get("brand")) or name.split()[0].title()
        sub_category = _cell_str(d.get("sub_category")).lower() or "food"
        category = _cell_str(d.get("category")).lower() or "catalogue"
        pet_type = _cell_str(d.get("pet_type")).lower() or "both"
        price = _cell_float(d.get("price"))
        status = _cell_str(d.get("status")).lower() or "published"
        image_url = _cell_str(d.get("image_url"))
        slug = _slugify(_cell_str(d.get("slug")) or name)

        docs.append({
            "id": str(uuid.uuid4()),
            "slug": slug,
            "name": name,
            "name_ka": _cell_str(d.get("name_ka")) or None,
            "brand": brand,
            "category": category,
            "sub_category": sub_category,
            "pet_type": pet_type,
            "size": _cell_str(d.get("size")) or None,
            "price": price,
            "currency": (_cell_str(d.get("currency")) or "GEL").upper(),
            "image": image_url or "https://images.unsplash.com/photo-1583337130417-3346a1be7dee?w=800",
            "description": _cell_str(d.get("description")) or f"{name} — please review.",
            "description_ka": _cell_str(d.get("description_ka")) or None,
            "tags": _cell_tags(d.get("tags")),
            "featured": _cell_bool(d.get("featured")),
            "status": status,
            "created_at": now_iso,
            "updated_at": now_iso,
        })

    # Dedupe by slug (later wins)
    by_slug = {d["slug"]: d for d in docs}
    docs = list(by_slug.values())
    print(f"Parsed {len(docs)} unique products from the file")

    client = AsyncIOMotorClient(MONGO_URL)
    db = client[DB_NAME]

    deleted = await db.products.delete_many({})
    print(f"Wiped {deleted.deleted_count} existing products")

    if docs:
        await db.products.insert_many(docs)

    total = await db.products.count_documents({})
    by_sc = {}
    async for d in db.products.find({}, {"sub_category": 1}):
        sc = d.get("sub_category", "unknown")
        by_sc[sc] = by_sc.get(sc, 0) + 1
    print(f"Inserted {total} products. By sub_category:")
    for k, v in sorted(by_sc.items()):
        print(f"  {k}: {v}")
    client.close()


if __name__ == "__main__":
    files = sorted(Path("/app/backend/uploads").glob("smartpaw-import-*.xlsx"))
    if not files:
        print("No scraped xlsx found in /app/backend/uploads")
        sys.exit(1)
    asyncio.run(main(files[-1]))
