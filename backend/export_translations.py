"""Export side-by-side English / Georgian UI strings to an Excel file.

Reads `frontend/src/lib/i18n/en.js` and `ka.js`, flattens every key into a
dotted path (e.g. ``footer.address``) and writes ``translation_review.xlsx``
to ``backend/uploads/public/`` so the client can download it from
``GET /api/uploads/public/translation_review.xlsx``.

Run with:
    python -m backend.export_translations          # from /app
or
    python /app/backend/export_translations.py
"""

from __future__ import annotations

import json
import re
import subprocess
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
EN_FILE = ROOT / "frontend" / "src" / "lib" / "i18n" / "en.js"
KA_FILE = ROOT / "frontend" / "src" / "lib" / "i18n" / "ka.js"
OUT_DIR = ROOT / "backend" / "uploads" / "public"
OUT_FILE = OUT_DIR / "translation_review.xlsx"


NODE_SCRIPT = r"""
import fs from 'node:fs';
import url from 'node:url';
const file = process.argv[2];
const varName = process.argv[3];
const tmp = file.replace(/\.js$/, '.__i18n_tmp__.mjs');
fs.copyFileSync(file, tmp);
try {
  const mod = await import(url.pathToFileURL(tmp).href);
  const data = mod[varName] || mod.default || mod;
  process.stdout.write(JSON.stringify(data));
} catch (e) {
  console.error(e);
  process.exitCode = 1;
} finally {
  try { fs.unlinkSync(tmp); } catch (_) {}
}
"""


def _load_with_node(js_file: Path, var_name: str) -> dict:
    """Use Node to evaluate the i18n JS module and return a plain dict."""
    script_path = OUT_DIR / "_load_i18n.mjs"
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    script_path.write_text(NODE_SCRIPT, encoding="utf-8")
    try:
        result = subprocess.run(
            ["node", str(script_path), str(js_file), var_name],
            check=True,
            capture_output=True,
            text=True,
        )
    finally:
        try:
            script_path.unlink()
        except OSError:
            pass
    return json.loads(result.stdout)


def _strip_js(src: str) -> str:
    src = re.sub(r"/\*.*?\*/", "", src, flags=re.S)
    src = re.sub(r"//.*", "", src)
    return src


def _load_with_regex(js_file: Path, var_name: str) -> dict:
    """Fallback parser when Node is unavailable.

    Converts the JS object literal into JSON by stripping comments,
    quoting unquoted keys and removing trailing commas.
    """
    src = js_file.read_text(encoding="utf-8")
    src = _strip_js(src)
    match = re.search(rf"const\s+{var_name}\s*=\s*(\{{)", src)
    if not match:
        raise RuntimeError(f"Could not find `const {var_name} =` in {js_file}")
    start = match.start(1)
    depth = 0
    end = None
    in_str: str | None = None
    escape = False
    for i, ch in enumerate(src[start:], start=start):
        if in_str:
            if escape:
                escape = False
            elif ch == "\\":
                escape = True
            elif ch == in_str:
                in_str = None
            continue
        if ch in ("'", '"', "`"):
            in_str = ch
            continue
        if ch == "{":
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0:
                end = i + 1
                break
    if end is None:
        raise RuntimeError(f"Unbalanced braces while parsing {js_file}")
    obj_src = src[start:end]
    obj_src = re.sub(r"([{,]\s*)([A-Za-z_][A-Za-z0-9_]*)\s*:", r'\1"\2":', obj_src)
    obj_src = obj_src.replace("'", '"')
    obj_src = re.sub(r",(\s*[}\]])", r"\1", obj_src)
    return json.loads(obj_src)


def load_dict(js_file: Path, var_name: str) -> dict:
    try:
        return _load_with_node(js_file, var_name)
    except (FileNotFoundError, subprocess.CalledProcessError):
        return _load_with_regex(js_file, var_name)


def flatten(prefix: str, value, out: list[tuple[str, str]]) -> None:
    if isinstance(value, dict):
        for k, v in value.items():
            key = f"{prefix}.{k}" if prefix else k
            flatten(key, v, out)
    elif isinstance(value, list):
        for idx, v in enumerate(value):
            flatten(f"{prefix}[{idx}]", v, out)
    else:
        out.append((prefix, "" if value is None else str(value)))


def build_rows() -> list[tuple[str, str, str]]:
    en = load_dict(EN_FILE, "en")
    ka = load_dict(KA_FILE, "ka")

    en_flat: list[tuple[str, str]] = []
    ka_flat: list[tuple[str, str]] = []
    flatten("", en, en_flat)
    flatten("", ka, ka_flat)
    ka_map = dict(ka_flat)

    rows: list[tuple[str, str, str]] = []
    for key, en_value in en_flat:
        rows.append((key, en_value, ka_map.get(key, "")))

    extra = [k for k, _ in ka_flat if k not in dict(en_flat)]
    for key in extra:
        rows.append((key, "", dict(ka_flat)[key]))

    return rows


def write_xlsx(rows: list[tuple[str, str, str]]) -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Translations"

    headers = ["Key", "English", "Georgian", "Notes"]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="0A4D8C")
    header_font = Font(bold=True, color="FFFFFF")
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")

    for key, en_value, ka_value in rows:
        ws.append([key, en_value, ka_value, ""])

    widths = [38, 70, 70, 30]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"
    wb.save(OUT_FILE)
    return OUT_FILE


def main() -> None:
    rows = build_rows()
    path = write_xlsx(rows)
    print(f"Wrote {len(rows)} translation rows to {path}")


if __name__ == "__main__":
    main()
