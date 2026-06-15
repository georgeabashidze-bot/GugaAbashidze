"""Build /app/frontend/src/lib/i18n.js from the revised translation Excel.

Reads /app/backend/uploads/imports/translation_review_revised.xlsx (sheet
'Translations' with columns Key | English | Georgian | Notes) and produces a
JS file that exports `translations = { en: {...}, ka: {...} }` plus the
existing LANGS array. Supports keys with array indices like
`faq.items[3].q` which build nested arrays.
"""

import json
import re
from pathlib import Path
from openpyxl import load_workbook

EXCEL = Path('/app/backend/uploads/imports/translation_review_revised.xlsx')
OUTPUT = Path('/app/frontend/src/lib/i18n.js')

# Legacy sections used by Breadcrumbs.jsx and PageShell.jsx — not present in
# the Excel sheet, so we preserve the prior Georgian translations here.
LEGACY = {
    'en': {
        'pageShell': {
            'comingSoon': 'Coming soon',
            'ctaTitle': 'Ready to skip the next pet-shop run?',
            'ctaStart': 'Start your plan',
            'ctaWhatsapp': 'Chat on WhatsApp',
        },
        'breadcrumbs': {
            'home': 'Home',
            'labels': {
                '/catalogue': 'Catalogue',
                '/catalogue/food': 'Food',
                '/catalogue/hygiene': 'Hygiene',
                '/catalogue/vitamins': 'Vitamins & Additives',
                '/special-offers': 'Special Offers',
                '/special-offers/toys-accessories': 'Toys & Accessories',
                '/special-offers/innovation-tech': 'Innovation & Tech',
                '/special-offers/services': 'Services',
                '/how-it-works': 'How It Works',
                '/plans': 'Plans & Pricing',
                '/about': 'About',
                '/blog': 'Blog',
                '/contact': 'Contact',
                '/faq': 'FAQ',
                '/privacy': 'Privacy Policy',
                '/terms': 'Terms & Conditions',
                '/delivery-policy': 'Delivery Policy',
                '/refund-policy': 'Refund Policy',
            },
        },
    },
    'ka': {
        'pageShell': {
            'comingSoon': 'მალე ხელმისაწვდომი',
            'ctaTitle': 'მზად ხარ შემდეგი ცხოველის მაღაზიის სვლის გამოტოვებას?',
            'ctaStart': 'დაიწყე გეგმა',
            'ctaWhatsapp': 'WhatsApp-ით საუბარი',
        },
        'breadcrumbs': {
            'home': 'მთავარი',
            'labels': {
                '/catalogue': 'კატალოგი',
                '/catalogue/food': 'საკვები',
                '/catalogue/hygiene': 'ჰიგიენა',
                '/catalogue/vitamins': 'ვიტამინები და დანამატები',
                '/special-offers': 'სპეციალური შეთავაზებები',
                '/special-offers/toys-accessories': 'სათამაშოები და აქსესუარები',
                '/special-offers/innovation-tech': 'ინოვაცია და ტექნოლოგია',
                '/special-offers/services': 'სერვისები',
                '/how-it-works': 'როგორ მუშაობს',
                '/plans': 'გეგმები და ფასები',
                '/about': 'ჩვენ შესახებ',
                '/blog': 'ბლოგი',
                '/contact': 'კონტაქტი',
                '/faq': 'ხშირი კითხვები',
                '/privacy': 'კონფიდენციალურობის პოლიტიკა',
                '/terms': 'წესები და პირობები',
                '/delivery-policy': 'მიწოდების პოლიტიკა',
                '/refund-policy': 'დაბრუნების პოლიტიკა',
            },
        },
    },
}

TOKEN_RE = re.compile(r'([^.\[\]]+)|\[(\d+)\]')


def parse_path(key):
    """'faq.items[3].q' -> [('k','faq'), ('k','items'), ('i',3), ('k','q')]"""
    tokens = []
    for m in TOKEN_RE.finditer(key):
        if m.group(1) is not None:
            tokens.append(('k', m.group(1)))
        else:
            tokens.append(('i', int(m.group(2))))
    return tokens


def set_path(root, tokens, value):
    cur = root
    for idx, (kind, tok) in enumerate(tokens):
        is_last = idx == len(tokens) - 1
        nxt_kind = None if is_last else tokens[idx + 1][0]
        if kind == 'k':
            if is_last:
                cur[tok] = value
            else:
                if tok not in cur or cur[tok] is None:
                    cur[tok] = [] if nxt_kind == 'i' else {}
                cur = cur[tok]
        else:  # index
            while len(cur) <= tok:
                cur.append({} if nxt_kind in ('k', None) else [])
            if is_last:
                cur[tok] = value
            else:
                if cur[tok] is None or (nxt_kind == 'i' and not isinstance(cur[tok], list)) or (nxt_kind == 'k' and not isinstance(cur[tok], dict)):
                    cur[tok] = [] if nxt_kind == 'i' else {}
                cur = cur[tok]


def build_dicts():
    wb = load_workbook(EXCEL, data_only=True)
    ws = wb['Translations']
    en, ka = {}, {}
    rows = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        key, en_val, ka_val, _notes = row[0], row[1], row[2], row[3] if len(row) > 3 else None
        key = str(key).strip()
        if not key:
            continue
        en_val = '' if en_val is None else str(en_val)
        ka_val = '' if ka_val is None else str(ka_val)
        # If the key is a single bare segment with no '.' / '[', store at root
        tokens = parse_path(key)
        set_path(en, tokens, en_val)
        set_path(ka, tokens, ka_val)
        rows += 1
    return en, ka, rows


def write_output(en, ka):
    en_json = json.dumps(en, ensure_ascii=False, indent=2)
    ka_json = json.dumps(ka, ensure_ascii=False, indent=2)
    content = (
        "// AUTO-GENERATED from /app/backend/uploads/imports/translation_review_revised.xlsx\n"
        "// Regenerate via: python3 /app/backend/scripts/apply_translations.py\n"
        "\n"
        "export const translations = {\n"
        f"  en: {en_json},\n"
        f"  ka: {ka_json},\n"
        "};\n"
        "\n"
        "export const LANGS = [\n"
        "  { code: 'en', label: 'EN' },\n"
        "  { code: 'ka', label: 'KA' },\n"
        "];\n"
    )
    OUTPUT.write_text(content, encoding='utf-8')


if __name__ == '__main__':
    en, ka, rows = build_dicts()
    # Merge legacy sections (Excel takes precedence if collision)
    for grp, val in LEGACY['en'].items():
        en.setdefault(grp, val)
    for grp, val in LEGACY['ka'].items():
        ka.setdefault(grp, val)
    write_output(en, ka)
    print(f'Wrote {OUTPUT} from {rows} rows')
    print(f'EN top-level keys: {sorted(en.keys())}')
    print(f'KA top-level keys: {sorted(ka.keys())}')
