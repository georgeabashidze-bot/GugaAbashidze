"""Generate the partner-brand intake spreadsheet for the catalogue.

Output:  /app/backend/uploads/public/smartpaw_partner_brands.xlsx
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path("/app/backend/uploads/public/smartpaw_partner_brands.xlsx")
OUT.parent.mkdir(parents=True, exist_ok=True)

HEADER_FILL = PatternFill("solid", fgColor="0A4D8C")
ACCENT_FILL = PatternFill("solid", fgColor="F5F2EB")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
BORDER = Border(
    left=Side(style="thin", color="DDDDDD"),
    right=Side(style="thin", color="DDDDDD"),
    top=Side(style="thin", color="DDDDDD"),
    bottom=Side(style="thin", color="DDDDDD"),
)
TOP_WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()

# ---------------------------------------------------------------------------
# Sheet 1: README / instructions
# ---------------------------------------------------------------------------
ws = wb.active
ws.title = "How to use"
ws.append(["#", "English", "ქართული"])
for c in ws[1]:
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
    c.alignment = Alignment(vertical="center")

steps = [
    ("1", "Fill in one row per partner brand on the 'Brands' sheet.",
           "გადადი ფურცელზე „Brands“ და შეავსე თითო მწკრივი თითო პარტნიორი ბრენდისთვის."),
    ("2", "The 'Category' column accepts: Food, Hygiene, Vitamins. A brand that sells in multiple categories must be added on a separate row for each category.",
           "სვეტი „Category“ მიიღებს: Food, Hygiene, Vitamins. თუ ბრენდი მუშაობს რამდენიმე კატეგორიაში — დაამატე ცალკე მწკრივი თითო კატეგორიისთვის."),
    ("3", "The 'Subtype' column applies ONLY to Food and Hygiene. For Vitamins, leave it blank.",
           "სვეტი „Subtype“ ეხება მხოლოდ Food-ს და Hygiene-ს. Vitamins-ში დატოვე ცარიელი."),
    ("4", "Allowed subtypes:\n  Food → Dry Food, Wet Food, Snacks, Other\n  Hygiene → Grooming, Veterinary Line, Specific Care, Other",
           "დასაშვები Subtype-ები:\n  Food → Dry Food, Wet Food, Snacks, Other\n  Hygiene → Grooming, Veterinary Line, Specific Care, Other"),
    ("5", "Logo URL: paste a direct image URL (PNG / SVG / JPG). If you don't have one yet, leave blank — a text-based placeholder will be used.",
           "Logo URL: ჩასვი ლოგოს პირდაპირი ბმული (PNG / SVG / JPG). თუ ჯერ არ გაქვს, დატოვე ცარიელი — გამოჩნდება ტექსტური ჩანაცვლება."),
    ("6", "Base URL: the partner's catalogue page (or homepage) — opens in a new tab.",
           "Base URL: პარტნიორის კატალოგის გვერდი (ან მთავარი) — გაიხსნება ახალ ფანჯარაში."),
    ("7", "Subtype URL (optional): a deep link to that brand's specific section, e.g. Royal Canin → Dry food. If you leave it empty, the Base URL is used.",
           "Subtype URL (არასავალდებულო): ღრმა ბმული ბრენდის შესაბამის სექციაში (მაგ. Royal Canin → Dry food). თუ ცარიელია — გამოყენებული იქნება Base URL."),
    ("8", "Once you've filled it in, send the file back to me. I'll wire it into the live catalogue.",
           "შევსების შემდეგ დააბრუნე ფაილი ჩემთან. გავაერთიანებ ცოცხალ კატალოგში."),
]

for s in steps:
    ws.append(s)

for col, w in enumerate([6, 64, 64], start=1):
    ws.column_dimensions[get_column_letter(col)].width = w
ws.row_dimensions[1].height = 26
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = TOP_WRAP
        cell.border = BORDER

# ---------------------------------------------------------------------------
# Sheet 2: Brands (the editable one)
# ---------------------------------------------------------------------------
ws = wb.create_sheet("Brands")
headers = [
    "Brand Name (EN)",
    "Brand Name (KA)",
    "Category",
    "Subtype",
    "Logo URL",
    "Base URL",
    "Subtype URL (optional)",
    "Notes",
]
ws.append(headers)
for c in ws[1]:
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
    c.alignment = Alignment(vertical="center", wrap_text=True)
ws.row_dimensions[1].height = 36

# Sample rows so the user has examples (formatted with accent fill)
samples = [
    ("Royal Canin",  "Royal Canin",  "Food",     "Dry Food",         "https://example.com/royalcanin.png",
        "https://royalcanin.ge", "https://royalcanin.ge/dry-food", "Example only — replace with your real brand list"),
    ("Hill's Science Diet", "Hill's", "Food",    "Wet Food",         "",
        "https://hillspet.com", "", "Example only"),
    ("Bewital", "Bewital", "Hygiene", "Grooming", "",
        "https://bewital.com", "", "Example only"),
    ("Virbac", "Virbac", "Vitamins", "", "",
        "https://virbac.com", "", "Example only — Vitamins has no Subtype"),
]
sample_start = 2
for s in samples:
    ws.append(s)
for r in range(sample_start, sample_start + len(samples)):
    for cell in ws[r]:
        cell.fill = ACCENT_FILL
        cell.alignment = TOP_WRAP
        cell.border = BORDER

# Blank editable rows (50 of them)
for _ in range(50):
    ws.append([""] * len(headers))
for row in ws.iter_rows(min_row=sample_start + len(samples)):
    for cell in row:
        cell.alignment = TOP_WRAP
        cell.border = BORDER

# Column widths
widths = [22, 22, 12, 22, 38, 38, 38, 30]
for col, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(col)].width = w

# Data validation for Category column (C)
cat_dv = DataValidation(
    type="list",
    formula1='"Food,Hygiene,Vitamins"',
    allow_blank=True,
    showDropDown=False,  # show the dropdown arrow
)
cat_dv.error = "Pick one of: Food, Hygiene, Vitamins"
cat_dv.errorTitle = "Invalid category"
cat_dv.add(f"C2:C{ws.max_row}")
ws.add_data_validation(cat_dv)

# Data validation for Subtype column (D)
sub_dv = DataValidation(
    type="list",
    formula1='"Dry Food,Wet Food,Snacks,Other,Grooming,Veterinary Line,Specific Care"',
    allow_blank=True,
    showDropDown=False,
)
sub_dv.error = "Pick a valid subtype for Food or Hygiene. Leave blank for Vitamins."
sub_dv.errorTitle = "Invalid subtype"
sub_dv.add(f"D2:D{ws.max_row}")
ws.add_data_validation(sub_dv)

ws.freeze_panes = "A2"

# ---------------------------------------------------------------------------
# Sheet 3: Reference (allowed values)
# ---------------------------------------------------------------------------
ws = wb.create_sheet("Reference")
ws.append(["Category", "Allowed Subtype values"])
for c in ws[1]:
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
ref_rows = [
    ("Food",     "Dry Food, Wet Food, Snacks, Other"),
    ("Hygiene",  "Grooming, Veterinary Line, Specific Care, Other"),
    ("Vitamins", "(no subtype — leave blank)"),
]
for r in ref_rows:
    ws.append(r)
for col, w in enumerate([14, 54], start=1):
    ws.column_dimensions[get_column_letter(col)].width = w
for row in ws.iter_rows(min_row=1):
    for cell in row:
        cell.alignment = TOP_WRAP
        cell.border = BORDER

wb.save(OUT)
print(f"Saved → {OUT}")
print(f"Size : {OUT.stat().st_size:,} bytes")
