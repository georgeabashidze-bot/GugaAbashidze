"""Build a bulk-upload Excel for the Monge / Simba / LeChat catalogue
extracted from the supplier price-list (Monge MS Georgia, 01.10.25).

Reads the parsed JSON in memory below, normalises every row to the
SmartPaw bulk-import template schema (see backend/products_import.py)
and writes the file to ``backend/uploads/public/monge-catalogue.xlsx``
so the client can download it from
``GET /api/uploads/public/monge-catalogue.xlsx``.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT_DIR = Path(__file__).resolve().parent.parent / "uploads" / "public"
OUT_FILE = OUT_DIR / "monge-catalogue.xlsx"

# ----------------------------------------------------------------------
# Source rows parsed from the supplier PDF (Monge 01.10.25 retail prices)
# ----------------------------------------------------------------------
SOURCE: list[dict] = json.loads(r"""
[
{"code":"4008","name":"MONGE Natural Superpremium Mini Starter Rich in Chicken","brand":"Monge","pet_type":"dog","life_stage":"puppy","form":"dry-food","size":"1.5 kg","price":37.75},
{"code":"4015","name":"MONGE Natural Superpremium Mini Puppy & Junior Rich in Chicken","brand":"Monge","pet_type":"dog","life_stage":"puppy","form":"dry-food","size":"3 kg","price":72.90},
{"code":"4114","name":"MONGE Natural Superpremium Mini Adult Rich in Chicken","brand":"Monge","pet_type":"dog","life_stage":"adult","form":"dry-food","size":"3 kg","price":54.15},
{"code":"4053","name":"MONGE Natural Superpremium Medium Starter Rich in Chicken","brand":"Monge","pet_type":"dog","life_stage":"puppy","form":"dry-food","size":"7.5 kg","price":135.40},
{"code":"1600","name":"MONGE Natural Superpremium Medium Starter Rich in Chicken","brand":"Monge","pet_type":"dog","life_stage":"puppy","form":"dry-food","size":"1.5 kg","price":35.50},
{"code":"1662","name":"MONGE Natural Superpremium Medium Puppy & Junior Rich in Chicken","brand":"Monge","pet_type":"dog","life_stage":"puppy","form":"dry-food","size":"12 kg","price":215.90},
{"code":"4275","name":"MONGE Natural Superpremium Medium Adult Rich in Chicken","brand":"Monge","pet_type":"dog","life_stage":"adult","form":"dry-food","size":"12 kg","price":195.35},
{"code":"1709","name":"MONGE Natural Superpremium Maxi Puppy & Junior Rich in Chicken","brand":"Monge","pet_type":"dog","life_stage":"puppy","form":"dry-food","size":"15 kg","price":218.10,"note":"PDF showed '215' — assumed 15 kg (Maxi line). Please confirm."},
{"code":"4411","name":"MONGE Natural Superpremium Maxi Adult Rich in Chicken","brand":"Monge","pet_type":"dog","life_stage":"adult","form":"dry-food","size":"12 kg","price":191.50},
{"code":"1525","name":"MONGE Natural Superpremium Mini Puppy & Junior Monoprotein Lamb with Rice","brand":"Monge","pet_type":"dog","life_stage":"puppy","form":"dry-food","size":"2.5 kg","price":64.95},
{"code":"1075","name":"MONGE Natural Superpremium Mini Adult Monoprotein Lamb with Rice and Potatoes","brand":"Monge","pet_type":"dog","life_stage":"adult","form":"dry-food","size":"7.5 kg","price":165.70},
{"code":"1556","name":"MONGE Natural Superpremium Mini Adult Monoprotein Lamb with Rice and Potatoes","brand":"Monge","pet_type":"dog","life_stage":"adult","form":"dry-food","size":"7.5 kg","price":160.40,"note":"PDF appears to duplicate row 11 with code 1556 — please confirm if this is a different variant."},
{"code":"1051","name":"MONGE Natural Superpremium Mini Puppy & Junior Monoprotein Salmon with Rice","brand":"Monge","pet_type":"dog","life_stage":"puppy","form":"dry-food","size":"2.5 kg","price":65.50},
{"code":"1068","name":"MONGE Natural Superpremium Mini Adult Monoprotein Salmon with Rice","brand":"Monge","pet_type":"dog","life_stage":"adult","form":"dry-food","size":"7.5 kg","price":166.00},
{"code":"1570","name":"MONGE Natural Superpremium Mini Adult Monoprotein Salmon with Rice","brand":"Monge","pet_type":"dog","life_stage":"adult","form":"dry-food","size":"2.5 kg","price":62.40},
{"code":"1587","name":"MONGE Natural Superpremium Mini Adult Monoprotein Salmon with Rice","brand":"Monge","pet_type":"dog","life_stage":"adult","form":"dry-food","size":"7.5 kg","price":161.50,"note":"PDF appears to duplicate row 14 with code 1587 — please confirm."},
{"code":"1198","name":"MONGE Natural Superpremium All Breeds Puppy & Junior Monoprotein Lamb with Rice","brand":"Monge","pet_type":"dog","life_stage":"puppy","form":"dry-food","size":"12 kg","price":251.95},
{"code":"1327","name":"MONGE Natural Superpremium All Breeds Adult Monoprotein Lamb with Rice and Potatoes","brand":"Monge","pet_type":"dog","life_stage":"adult","form":"dry-food","size":"12 kg","price":237.05},
{"code":"1211","name":"MONGE Natural Superpremium All Breeds Puppy & Junior Monoprotein Salmon with Rice","brand":"Monge","pet_type":"dog","life_stage":"puppy","form":"dry-food","size":"12 kg","price":250.00},
{"code":"1303","name":"MONGE Natural Superpremium All Breeds Adult Monoprotein Salmon with Rice","brand":"Monge","pet_type":"dog","life_stage":"adult","form":"dry-food","size":"12 kg","price":244.70},
{"code":"1174","name":"MONGE Natural Superpremium All Breeds Adult Hypo with Salmon and Tuna","brand":"Monge","pet_type":"dog","life_stage":"adult","form":"dry-food","size":"12 kg","price":258.50},
{"code":"1365","name":"MONGE Natural Superpremium All Breeds Puppy & Junior Monoprotein Beef with Rice","brand":"Monge","pet_type":"dog","life_stage":"puppy","form":"dry-food","size":"12 kg","price":238.00},
{"code":"1341","name":"MONGE Natural Superpremium All Breeds Adult Monoprotein Beef with Rice","brand":"Monge","pet_type":"dog","life_stage":"adult","form":"dry-food","size":"12 kg","price":236.00},
{"code":"1396","name":"MONGE Natural Superpremium All Breeds Adult Monoprotein Turkey with Rice and Potatoes","brand":"Monge","pet_type":"dog","life_stage":"adult","form":"dry-food","size":"12 kg","price":244.70},
{"code":"1136","name":"MONGE Natural Superpremium All Breeds Adult Monoprotein Duck with Rice and Potatoes","brand":"Monge","pet_type":"dog","life_stage":"adult","form":"dry-food","size":"12 kg","price":245.30},
{"code":"4441","name":"Monge Fresh Puppy Dog Chunks in Loaf Veal with Vegetables","brand":"Monge Fresh","pet_type":"dog","life_stage":"puppy","form":"wet-food","size":"400 g","price":6.15},
{"code":"4458","name":"Monge Fresh Adult Dog Chunks in Loaf with Veal","brand":"Monge Fresh","pet_type":"dog","life_stage":"adult","form":"wet-food","size":"400 g","price":6.00},
{"code":"4571","name":"Monge Fresh Adult Dog Chunks in Loaf with Lamb","brand":"Monge Fresh","pet_type":"dog","life_stage":"adult","form":"wet-food","size":"400 g","price":6.00},
{"code":"4564","name":"Monge Fresh Adult Dog Chunks in Loaf with Duck","brand":"Monge Fresh","pet_type":"dog","life_stage":"adult","form":"wet-food","size":"400 g","price":6.00},
{"code":"4472","name":"Monge Fresh Adult Dog Chunks in Loaf with Chicken","brand":"Monge Fresh","pet_type":"dog","life_stage":"adult","form":"wet-food","size":"400 g","price":6.00},
{"code":"2037","name":"Monge VetSolution Canine Hypo Monoprotein Tuna","brand":"Monge VetSolution","pet_type":"dog","life_stage":"adult","form":"wet-food","size":"400 g","price":9.75,"note":"VetSolution Hypo wet pouch — please confirm form (wet-food)."},
{"code":"2051","name":"Monge VetSolution Canine Hypo Monoprotein Lamb","brand":"Monge VetSolution","pet_type":"dog","life_stage":"adult","form":"wet-food","size":"400 g","price":9.75},
{"code":"2044","name":"Monge VetSolution Canine Hypo Monoprotein Duck","brand":"Monge VetSolution","pet_type":"dog","life_stage":"adult","form":"wet-food","size":"400 g","price":9.75},
{"code":"3017","name":"MONGE Paté and Chunkies with Tuna","brand":"Monge","pet_type":"dog","life_stage":"all","form":"wet-food","size":"100 g","price":3.75},
{"code":"3024","name":"MONGE Paté and Chunkies with Turkey","brand":"Monge","pet_type":"dog","life_stage":"all","form":"wet-food","size":"100 g","price":3.75},
{"code":"3048","name":"MONGE Paté and Chunkies with Duck","brand":"Monge","pet_type":"dog","life_stage":"all","form":"wet-food","size":"100 g","price":3.75},
{"code":"3055","name":"MONGE Paté and Chunkies with Lamb","brand":"Monge","pet_type":"dog","life_stage":"all","form":"wet-food","size":"100 g","price":3.75},
{"code":"3062","name":"MONGE Paté and Chunkies with Chicken","brand":"Monge","pet_type":"dog","life_stage":"all","form":"wet-food","size":"100 g","price":3.75},
{"code":"3079","name":"MONGE Paté and Chunkies with Beef","brand":"Monge","pet_type":"dog","life_stage":"all","form":"wet-food","size":"100 g","price":3.75},
{"code":"3208","name":"MONGE Paté and Chunkies with Turkey and Blueberry","brand":"Monge","pet_type":"dog","life_stage":"all","form":"wet-food","size":"100 g","price":3.50},
{"code":"3215","name":"MONGE Paté and Chunkies with Chicken and Raspberry","brand":"Monge","pet_type":"dog","life_stage":"all","form":"wet-food","size":"100 g","price":3.50},
{"code":"3222","name":"MONGE Paté and Chunkies with Lamb and Apple","brand":"Monge","pet_type":"dog","life_stage":"all","form":"wet-food","size":"100 g","price":3.50},
{"code":"3239","name":"MONGE Paté and Chunkies with Duck and Orange","brand":"Monge","pet_type":"dog","life_stage":"all","form":"wet-food","size":"100 g","price":3.50},
{"code":"3246","name":"MONGE Paté and Chunkies with Salmon and Pear","brand":"Monge","pet_type":"dog","life_stage":"all","form":"wet-food","size":"100 g","price":3.50},
{"code":"6243","name":"MONGE Natural Superpremium Kitten Rich in Chicken","brand":"Monge","pet_type":"cat","life_stage":"kitten","form":"dry-food","size":"10 kg","price":237.05},
{"code":"6236","name":"MONGE Natural Superpremium Adult Rich in Chicken","brand":"Monge","pet_type":"cat","life_stage":"adult","form":"dry-food","size":"10 kg","price":207.00},
{"code":"1914","name":"MONGE Natural Superpremium Urinary Feline Rich in Chicken","brand":"Monge","pet_type":"cat","life_stage":"adult","form":"dry-food","size":"1.5 kg","price":48.40},
{"code":"6267","name":"MONGE Natural Superpremium Sterilised Rich in Chicken","brand":"Monge","pet_type":"cat","life_stage":"adult","form":"dry-food","size":"10 kg","price":231.10},
{"code":"5500","name":"MONGE Natural Superpremium Monoprotein Kitten Cat Trout","brand":"Monge","pet_type":"cat","life_stage":"kitten","form":"dry-food","size":"1.5 kg","price":51.00},
{"code":"1976","name":"MONGE Natural Superpremium Monoprotein Kitten Beef","brand":"Monge","pet_type":"cat","life_stage":"kitten","form":"dry-food","size":"1.5 kg","price":52.00},
{"code":"5524","name":"MONGE Natural Superpremium Monoprotein Sterilised Cat Beef","brand":"Monge","pet_type":"cat","life_stage":"adult","form":"dry-food","size":"1.5 kg","price":50.00},
{"code":"6304","name":"MONGE Natural Superpremium Monoprotein Sterilised Cat Trout","brand":"Monge","pet_type":"cat","life_stage":"adult","form":"dry-food","size":"10 kg","price":243.00},
{"code":"6182","name":"MONGE Natural Superpremium Monoprotein Sterilised Cat Duck","brand":"Monge","pet_type":"cat","life_stage":"adult","form":"dry-food","size":"10 kg","price":241.50},
{"code":"1945","name":"MONGE Natural Superpremium Monoprotein Adult Cat Rabbit","brand":"Monge","pet_type":"cat","life_stage":"adult","form":"dry-food","size":"1.5 kg","price":50.60},
{"code":"6298","name":"MONGE Natural Superpremium Monoprotein Adult Cat Salmon","brand":"Monge","pet_type":"cat","life_stage":"adult","form":"dry-food","size":"10 kg","price":222.50},
{"code":"3604","name":"MONGE Grill Kitten Cat Chunkies in Jelly Rich in Salmon","brand":"Monge Grill","pet_type":"cat","life_stage":"kitten","form":"wet-food","size":"85 g","price":2.75},
{"code":"3611","name":"MONGE Grill Adult Cat Chunkies in Jelly Rich in Rabbit","brand":"Monge Grill","pet_type":"cat","life_stage":"adult","form":"wet-food","size":"85 g","price":2.75},
{"code":"3628","name":"MONGE Grill Adult Cat Chunkies in Jelly Rich in Lamb","brand":"Monge Grill","pet_type":"cat","life_stage":"adult","form":"wet-food","size":"85 g","price":2.75},
{"code":"3635","name":"MONGE Grill Sterilised Cat Chunkies in Jelly Rich in Cockerel","brand":"Monge Grill","pet_type":"cat","life_stage":"adult","form":"wet-food","size":"85 g","price":2.75},
{"code":"3642","name":"MONGE Grill Sterilised Cat Chunkies in Jelly Rich in Veal","brand":"Monge Grill","pet_type":"cat","life_stage":"adult","form":"wet-food","size":"85 g","price":2.75},
{"code":"3659","name":"MONGE Grill Sterilised Cat Chunkies in Jelly Rich in Trout","brand":"Monge Grill","pet_type":"cat","life_stage":"adult","form":"wet-food","size":"85 g","price":2.75},
{"code":"3727","name":"MONGE Solo Cat Monoprotein Paté Kitten Trout","brand":"Monge Solo","pet_type":"cat","life_stage":"kitten","form":"wet-food","size":"85 g","price":2.65},
{"code":"3703","name":"MONGE Solo Cat Monoprotein Paté Adult Duck","brand":"Monge Solo","pet_type":"cat","life_stage":"adult","form":"wet-food","size":"85 g","price":2.65},
{"code":"3734","name":"MONGE Solo Cat Monoprotein Paté Adult Turkey","brand":"Monge Solo","pet_type":"cat","life_stage":"adult","form":"wet-food","size":"85 g","price":2.65},
{"code":"3741","name":"MONGE Solo Cat Monoprotein Paté Sterilised Beef","brand":"Monge Solo","pet_type":"cat","life_stage":"adult","form":"wet-food","size":"85 g","price":2.65},
{"code":"3710","name":"MONGE Solo Cat Monoprotein Paté Sterilised Chicken","brand":"Monge Solo","pet_type":"cat","life_stage":"adult","form":"wet-food","size":"85 g","price":2.65},
{"code":"4533","name":"MONGE VetSolution Recovery (Canine)","brand":"Monge VetSolution","pet_type":"dog","life_stage":"all","form":"wet-food","size":"150 g","price":4.40},
{"code":"4526","name":"MONGE VetSolution Renal and Oxalate (Canine)","brand":"Monge VetSolution","pet_type":"dog","life_stage":"all","form":"wet-food","size":"150 g","price":4.40},
{"code":"4519","name":"MONGE VetSolution Gastrointestinal (Canine)","brand":"Monge VetSolution","pet_type":"dog","life_stage":"all","form":"wet-food","size":"150 g","price":4.40},
{"code":"4502","name":"MONGE VetSolution Dermatosis (Canine)","brand":"Monge VetSolution","pet_type":"dog","life_stage":"all","form":"wet-food","size":"150 g","price":4.35},
{"code":"1146","name":"MONGE Dry VetSolution Renal and Oxalate (Canine)","brand":"Monge VetSolution","pet_type":"dog","life_stage":"all","form":"dry-food","size":"12 kg","price":324.35},
{"code":"1047","name":"MONGE Dry VetSolution Gastrointestinal Adult (Canine)","brand":"Monge VetSolution","pet_type":"dog","life_stage":"adult","form":"dry-food","size":"2 kg","price":65.00},
{"code":"1054","name":"MONGE Dry VetSolution Dermatosis (Canine)","brand":"Monge VetSolution","pet_type":"dog","life_stage":"all","form":"dry-food","size":"12 kg","price":302.10},
{"code":"1016","name":"MONGE Dry VetSolution Dermatosis (Canine)","brand":"Monge VetSolution","pet_type":"dog","life_stage":"all","form":"dry-food","size":"12 kg","price":356.25,"note":"PDF lists two Dermatosis dry SKUs (1054 & 1016) at different prices — please confirm which variant."},
{"code":"4625","name":"MONGE VetSolution Urinary Struvite (Feline)","brand":"Monge VetSolution","pet_type":"cat","life_stage":"all","form":"wet-food","size":"100 g","price":3.65},
{"code":"4649","name":"MONGE VetSolution Renal and Oxalate (Feline)","brand":"Monge VetSolution","pet_type":"cat","life_stage":"all","form":"wet-food","size":"100 g","price":3.60},
{"code":"4618","name":"MONGE VetSolution Gastrointestinal (Feline)","brand":"Monge VetSolution","pet_type":"cat","life_stage":"all","form":"wet-food","size":"100 g","price":3.70},
{"code":"4601","name":"MONGE VetSolution Dermatosis (Feline)","brand":"Monge VetSolution","pet_type":"cat","life_stage":"all","form":"wet-food","size":"100 g","price":3.65},
{"code":"4656","name":"MONGE VetSolution Recovery (Feline)","brand":"Monge VetSolution","pet_type":"cat","life_stage":"all","form":"wet-food","size":"100 g","price":3.70},
{"code":"1580","name":"MONGE Dry VetSolution Urinary Struvite (Feline)","brand":"Monge VetSolution","pet_type":"cat","life_stage":"adult","form":"dry-food","size":"1.5 kg","price":65.50},
{"code":"1511","name":"MONGE Dry VetSolution Gastrointestinal (Feline)","brand":"Monge VetSolution","pet_type":"cat","life_stage":"adult","form":"dry-food","size":"1.5 kg","price":68.90},
{"code":"1658","name":"MONGE Dry VetSolution Renal (Feline)","brand":"Monge VetSolution","pet_type":"cat","life_stage":"adult","form":"dry-food","size":"1.5 kg","price":68.90},
{"code":"9867","name":"SIMBA Croquettes with Beef","brand":"Simba","pet_type":"dog","life_stage":"all","form":"dry-food","size":"20 kg","price":138.55},
{"code":"9874","name":"SIMBA Croquettes with Chicken","brand":"Simba","pet_type":"dog","life_stage":"all","form":"dry-food","size":"20 kg","price":138.55},
{"code":"9027","name":"SIMBA Chunks with Chicken and Turkey","brand":"Simba","pet_type":"dog","life_stage":"all","form":"wet-food","size":"415 g","price":3.60},
{"code":"9164","name":"SIMBA Chunks with Lamb","brand":"Simba","pet_type":"dog","life_stage":"all","form":"wet-food","size":"415 g","price":3.60},
{"code":"9171","name":"SIMBA Chunks with Wild Game","brand":"Simba","pet_type":"dog","life_stage":"all","form":"wet-food","size":"415 g","price":3.60},
{"code":"8600","name":"LeChat Excellence Paté and Chunkies with Beef","brand":"LeChat","pet_type":"cat","life_stage":"all","form":"wet-food","size":"100 g","price":3.40},
{"code":"8617","name":"LeChat Excellence Paté and Chunkies with Chicken and Vegetable","brand":"LeChat","pet_type":"cat","life_stage":"all","form":"wet-food","size":"100 g","price":3.40},
{"code":"8624","name":"LeChat Excellence Paté and Chunkies with Duck","brand":"LeChat","pet_type":"cat","life_stage":"all","form":"wet-food","size":"100 g","price":3.40},
{"code":"8631","name":"LeChat Excellence Paté and Chunkies with Salmon","brand":"LeChat","pet_type":"cat","life_stage":"all","form":"wet-food","size":"100 g","price":3.40}
]
""")


# ----------------------------------------------------------------------
# Build composed English description + tag list per row
# ----------------------------------------------------------------------
def build_description(row: dict) -> str:
    pet = "dogs" if row["pet_type"] == "dog" else "cats"
    stage_map = {
        "puppy": "puppies",
        "kitten": "kittens",
        "adult": "adult " + pet,
        "senior": "senior " + pet,
        "all": pet + " of all life stages",
    }
    stage = stage_map.get(row["life_stage"], "all " + pet)
    form_label = "Dry kibble" if row["form"] == "dry-food" else ("Wet food" if row["form"] == "wet-food" else "Treat")
    # Extract flavour / protein hint from the name
    flavour = ""
    m = re.search(r"(?:with|Rich in|Monoprotein)\s+(.+)$", row["name"])
    if m:
        flavour = " — " + m.group(1).strip()
    return f"{form_label} for {stage}. {row['brand']} {row['size']} pack.{flavour}".strip()


def build_tags(row: dict) -> str:
    tags = [row["brand"].lower().replace(" ", "-"), row["pet_type"], row["form"]]
    if row["life_stage"] not in ("", "all"):
        tags.append(row["life_stage"])
    if "Monoprotein" in row["name"]:
        tags.append("monoprotein")
    if "VetSolution" in row["name"] or "VetSolution" in row["brand"]:
        tags.append("veterinary-diet")
    if "Sterilised" in row["name"]:
        tags.append("sterilised")
    if "Urinary" in row["name"]:
        tags.append("urinary")
    if "Hypo" in row["name"]:
        tags.append("hypoallergenic")
    # de-dupe preserving order
    seen, out = set(), []
    for t in tags:
        if t and t not in seen:
            seen.add(t); out.append(t)
    return ", ".join(out)


def build_slug(row: dict) -> str:
    base = re.sub(r"[^a-z0-9]+", "-", row["name"].lower()).strip("-")
    return f"{base}-{row['code']}"


# ----------------------------------------------------------------------
# Template schema MUST match products_import.py TEMPLATE_COLUMNS
# ----------------------------------------------------------------------
HEADERS = [
    "name", "name_ka", "brand", "category", "sub_category", "product_type",
    "pet_type", "size", "price", "currency", "image_url", "description",
    "description_ka", "tags", "featured", "status", "slug",
]

CATEGORY_VALUES = ("catalogue", "specials")
SUB_CATEGORY_VALUES = ("food", "hygiene", "vitamins", "toys-accessories", "innovation-tech", "services")
PRODUCT_TYPE_VALUES = ("dry-food", "wet-food", "snacks", "other", "teeth-care", "grooming", "pads")
PET_TYPE_VALUES = ("dog", "cat", "both")
STATUS_VALUES = ("draft", "published")


def to_template_row(row: dict) -> dict:
    return {
        "name": row["name"],
        "name_ka": "",  # client to fill / translator
        "brand": row["brand"],
        "category": "catalogue",
        "sub_category": "food",
        "product_type": row["form"],
        "pet_type": row["pet_type"],
        "size": row["size"],
        "price": row["price"],
        "currency": "GEL",
        "image_url": "",  # client to fill
        "description": build_description(row),
        "description_ka": "",
        "tags": build_tags(row),
        "featured": "FALSE",
        "status": "draft",  # safer default — published once approved
        "slug": build_slug(row),
    }


def write_workbook() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"

    header_fill = PatternFill("solid", start_color="FF0A4D8C", end_color="FF0A4D8C")
    header_font = Font(bold=True, color="FFFFFFFF", size=11)
    note_fill = PatternFill("solid", start_color="FFFFF6CC", end_color="FFFFF6CC")

    # Row 1: headers + supplier_code + note columns appended for review
    full_headers = HEADERS + ["supplier_code", "notes_for_smartpaw"]
    for i, h in enumerate(full_headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="left", vertical="center")

    widths = [42, 32, 18, 12, 14, 12, 10, 12, 10, 10, 30, 60, 30, 30, 10, 12, 32, 14, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

    # Data rows
    for src in SOURCE:
        tpl = to_template_row(src)
        out = [tpl[h] for h in HEADERS] + [src["code"], src.get("note", "")]
        ws.append(out)

    # Wrap-text + top align all data cells
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        # highlight the notes column if not empty
        note_cell = row_cells[-1]
        if note_cell.value:
            note_cell.fill = note_fill

    # Data validations
    def add_dv(values, col_idx):
        dv = DataValidation(
            type="list",
            formula1='"' + ",".join(values) + '"',
            allow_blank=True,
            showErrorMessage=True,
        )
        ws.add_data_validation(dv)
        col = get_column_letter(col_idx)
        dv.add(f"{col}2:{col}1000")

    add_dv(CATEGORY_VALUES, HEADERS.index("category") + 1)
    add_dv(SUB_CATEGORY_VALUES, HEADERS.index("sub_category") + 1)
    add_dv(PRODUCT_TYPE_VALUES, HEADERS.index("product_type") + 1)
    add_dv(PET_TYPE_VALUES, HEADERS.index("pet_type") + 1)
    add_dv(STATUS_VALUES, HEADERS.index("status") + 1)
    add_dv(("TRUE", "FALSE"), HEADERS.index("featured") + 1)

    # Field-guide sheet
    guide = wb.create_sheet("How to fill")
    guide.column_dimensions["A"].width = 26
    guide.column_dimensions["B"].width = 100
    rows = [
        ("Source", "Monge MS Georgia retail pricelist, 01.10.2025 (PDF supplied by client)."),
        ("Rows", f"{len(SOURCE)} products."),
        ("Categorisation", "All rows set to category=catalogue, sub_category=food, status=draft."),
        ("name_ka", "EMPTY — please fill in the Georgian product name (the pricelist itself only carries English Latin names)."),
        ("description_ka", "EMPTY — please translate or rewrite the English description in Georgian."),
        ("image_url", "EMPTY — please paste the public image URL for each SKU. We'll download & host the file once you upload."),
        ("supplier_code", "Internal Monge article code, kept for your reference. NOT imported into the site."),
        ("notes_for_smartpaw", "Yellow rows flag items that need confirmation (OCR ambiguities or duplicate codes in the PDF)."),
        ("status", "Set to 'draft' so nothing goes live until you publish. Change to 'published' once images are added and rows are reviewed."),
        ("Re-import behaviour", "Rows are upserted by slug. Re-uploading the file UPDATES the same products — safe to iterate."),
    ]
    guide.cell(row=1, column=1, value="Topic").font = Font(bold=True)
    guide.cell(row=1, column=2, value="Notes").font = Font(bold=True)
    for i, (k, v) in enumerate(rows, start=2):
        guide.cell(row=i, column=1, value=k).font = Font(bold=True)
        guide.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(OUT_FILE)
    return OUT_FILE


if __name__ == "__main__":
    path = write_workbook()
    print(f"Wrote {len(SOURCE)} products to {path}")
